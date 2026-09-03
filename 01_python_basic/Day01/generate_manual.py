from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ========== 标题 ==========
title = doc.add_heading('AI-Agent 开发必备 Python 知识点精讲', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

subtitle = doc.add_paragraph('副标题：只学有用的，摒弃冗余，直击 ReAct / Function Calling 痛点')
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
subtitle.runs[0].font.size = Pt(16)
subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)

# ========== 第一部分：核心知识点清单 ==========
doc.add_heading('第一部分：核心知识点清单（共 10 大模块）', level=1)

modules = [
    ("模块一：基础语法（地基）",
     "1. 变量与 f-string：name = \"GPT\" ; prompt = f\"你是{name}, 当前时间{datetime.now()}\"（拼Prompt的绝对主力，禁止用+硬拼）。\n"
     "2. List 与 Dict 存取：messages.append({\"role\": \"user\", \"content\": msg})；工具参数用 params.get(\"city\", \"默认值\") 防止KeyError。\n"
     "3. 条件判断与循环：if \"action\" in resp:；while step < max_steps: 控制Agent迭代次数。\n"
     "4. ⭐ 循环终止术（必补）：必须熟练使用 break 跳出死循环，搭配 continue 跳过错误轮次。"),
    
    ("模块二：字符串清洗（Agent 存活率核心）",
     "1. strip / split：raw.strip('`').strip() 去除 LLM 输出的 Markdown 反引号。\n"
     "2. removeprefix / removesuffix：去除 \"Action: \" 等多余前缀。\n"
     "3. 暴力提取 JSON（备胎技能）：re.search(r'\\{.*\\}', text).group() 应对模型不听话时强行捞取。"),
    
    ("模块三：函数与数据返回",
     "1. 定义工具函数：def search_weather(city: str) -> dict: 必须返回字典结构。\n"
     "2. 关键字参数传参：tool_executor(**tool_args) 解包字典为函数参数。"),
    
    ("模块四：异常处理（Agent 的救命稻草）",
     "1. try-except-else-finally 标准结构。\n"
     "2. 必须捕获的三大异常：json.JSONDecodeError、requests.exceptions.RequestException、KeyError/IndexError。"),
    
    ("模块五：文件与环境变量（安全底线）",
     "1. 上下文管理器：强制使用 with open(\"file.txt\", \"r\", encoding=\"utf-8\") as f: 自动释放资源。\n"
     "2. dotenv 读取密钥：load_dotenv(); API_KEY = os.getenv(\"OPENAI_KEY\")，严禁硬编码在代码里。"),
    
    ("模块六：JSON 序列化与反序列化",
     "1. json.loads(resp_str) 转字典（解析模型输出）。\n"
     "2. json.dumps(tool_result, ensure_ascii=False) 转字符串（传给模型时必加 ensure_ascii=False 防中文乱码）。"),
    
    ("模块七：环境与依赖管理",
     "1. python -m venv venv 建虚拟环境。\n"
     "2. pip freeze > requirements.txt 导出依赖。\n"
     "3. import 顺序规范：标准库 → 第三方库 → 自定义模块。"),
    
    ("模块八：第三方库必会三件套",
     "1. requests：requests.post(url, headers, json=data, timeout=30)。\n"
     "2. python-dotenv：管理 .env。\n"
     "3. pydantic：配合 BaseModel 做工具调用的入参校验（防注入脏数据）。"),
    
    ("模块九：时间处理（给 AI 植入时钟）",
     "from datetime import datetime 将当前时间注入 System Prompt，否则大模型永远活在 2021 年。"),
    
    ("模块十：Git 版本控制（职业习惯）",
     "git add . -> git commit -m \"feat: add tool call\" -> git push。.gitignore 必须包含：venv/、.env、*.pyc、__pycache__/。")
]

for heading, content in modules:
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph(content)
    p.paragraph_format.left_indent = Inches(0.3)

# ========== 第二部分：14天极速学习路径 ==========
doc.add_heading('第二部分：14 天极速学习路径与每日实操计划', level=1)

phases = [
    ("阶段一：语法热身（Day 1-3）——只看不练假把式",
     "• Day 1：掌握 f-string、list/dict 增删改查，手写一个 messages 消息队列模拟。\n"
     "• Day 2：练习 while 循环 + break，写一个“猜数字”游戏模拟 Agent 最大迭代。\n"
     "• Day 3：专攻异常捕获，故意写错 JSON 字符串，用 try-except 优雅处理。"),
    
    ("阶段二：核心交互能力（Day 4-7）——手写原生 ReAct 雏形",
     "• Day 4：用 requests 调用一个公共 API（如聚合数据天气），封装成 get_weather 函数。\n"
     "• Day 5：结合 json.loads 和 re，写一个解析函数，能把 LLM 返回的 Action Input 字符串安全转为字典。\n"
     "• Day 6：主循环搭建。在 while 里拼接 System Prompt + User Prompt，处理模型返回的 thought 和 action。\n"
     "• Day 7：整合 dotenv 读取真实 Key，跑通第一个“查天气-回复结果”的完整单轮 Agent。"),
    
    ("阶段三：工具扩展与健壮性（Day 8-11）——让 Agent 长出手脚",
     "• Day 8：注册 3 个工具（查天气、查时间、计算器），用 Dict 映射工具名到函数。\n"
     "• Day 9：引入 Pydantic 校验，确保模型传参符合 city: str、radius: int 等类型。\n"
     "• Day 10：增加 finally 日志写入，每次循环把 messages 长度和 Token 预估打印出来。\n"
     "• Day 11：配合 with open 实现对话历史持久化（存为 history.json），重启时自动加载。"),
    
    ("阶段四：调试与交付（Day 12-14）——稳定大于一切",
     "• Day 12：专门制造各种“脏数据”（加前缀、多换行、少括号），用 .strip() 和正则暴力清洗。\n"
     "• Day 13：项目结构整理：区分 tools/、core/、prompts/ 文件夹，导出 requirements.txt。\n"
     "• Day 14：Git 提交所有代码，写一份详细的 README.md 说明如何启动。")
]

for heading, content in phases:
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph(content)
    p.paragraph_format.left_indent = Inches(0.3)

# ========== 第三部分：新手必踩的5个大坑（表格） ==========
doc.add_heading('第三部分：新手必踩的 5 个大坑与化解方法', level=1)

table = doc.add_table(rows=6, cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# 表头
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '常见错误现象'
hdr_cells[1].text = '根本原因'
hdr_cells[2].text = '解决方案'

data = [
    ("JSONDecodeError: Expecting value", "模型输出了自然语言解释或 Markdown 标记", "先 strip('`').strip()，再用正则提取 \\{.*\\}"),
    ("KeyError: 'tool_calls'", "模型未返回工具调用字段", "用 dict.get('tool_calls', []) 兜底"),
    ("程序跑着跑着卡住不动", "while 没加 break 或最大步数限制", "强制写 if step > max_iterations: break"),
    ("API 调用报错 ConnectionError", "网络波动或超时未设", "requests.post(timeout=10) 并在 except 里重试 2 次"),
    ("读取 .env 返回 None", "文件不在根目录或变量名前后有空格", "os.path.join(BASE_DIR, '.env') 固定路径")
]

for i, row_data in enumerate(data, start=1):
    row = table.rows[i].cells
    row[0].text = row_data[0]
    row[1].text = row_data[1]
    row[2].text = row_data[2]

# ========== 第四部分：终极方法论 ==========
doc.add_heading('第四部分：终极方法论（“三步迭代法”）', level=1)
method = doc.add_paragraph(
    "1. 硬编码跑通：先把固定字符串、固定参数写死，保证链路通。\n"
    "2. 动态化替换：把固定值改成 input() 或变量。\n"
    "3. 异常兜底：给每个关键 json.loads 和 requests.get 包一层 try。\n\n"
    "请记住：写 Agent 不是比谁代码优雅，而是比谁容错性强。把 30% 的时间花在异常处理和字符串清洗上，你的 Agent 稳定度会远超 80% 的初学者。"
)
method.paragraph_format.left_indent = Inches(0.3)

# ========== 保存 ==========
doc.save('AI-Agent_Python学习手册.docx')
print("✅ 完整版《AI-Agent_Python学习手册.docx》生成成功！")